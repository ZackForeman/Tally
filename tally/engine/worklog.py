"""Read the haulier's work log (CSV or XLSX) into jobs. Structured input, so no model.

One row per piece of work. Recognised columns (any order, case-insensitive;
only date, reg and customer are required):

    date, time, reg, driver, customer, ticket_no, type, from_site, to_site,
    material, net_t, miles, recycled, full_return, daywork_hours, time_in,
    time_out, chute, arrived, unload_start, internal

``type`` is load (default), return, daywork or waiting. Daywork (time on
site) is given as ``time_in`` and ``time_out``, or as ``daywork_hours``;
``chute`` is y for chute work. ``time_in``, ``time_out``, ``arrived`` and
``unload_start`` are HH:MM on the row's date (or full date-times); a time
out earlier than the time in is the next day.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from tally.engine.models import Daywork, Job, Leg, Source

TRUE = {"y", "yes", "true", "1", "x"}


class WorkLogError(ValueError):
    pass


def read_rows(data: bytes, name: str) -> list[dict[str, str]]:
    if name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "").strip() for h in rows[0]]
        return [
            {h: "" if v is None else _cell(v) for h, v in zip(header, row)}
            for row in rows[1:]
            if any(v is not None for v in row)
        ]
    text = data.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


def parse_work_log(data: bytes, name: str, doc_id: str) -> tuple[list[Job], list[str]]:
    """Jobs, plus one problem message per row that couldn't be read."""
    jobs, problems = [], []
    for n, raw in enumerate(read_rows(data, name), start=2):
        row = {k.strip().lower(): (v or "").strip() for k, v in raw.items() if k}
        try:
            jobs.append(_job(row, f"{doc_id}-{n}", name))
        except (WorkLogError, ValueError, InvalidOperation) as e:
            problems.append(f"row {n}: {e}")
    return jobs, problems


def _job(row: dict[str, str], job_id: str, name: str) -> Job:
    for col in ("date", "reg", "customer"):
        if not row.get(col):
            raise WorkLogError(f"no {col}")
    day = _date(row["date"])
    kind = (row.get("type") or "load").lower()
    if kind not in {"load", "return", "daywork", "waiting"}:
        raise WorkLogError(f"unknown type {kind!r}")

    daywork = None
    if kind == "daywork" or row.get("daywork_hours") or row.get("time_in"):
        daywork = _daywork(row, day)

    arrived = _moment(day, row.get("arrived"))
    unload_start = _moment(day, row.get("unload_start"))
    if arrived and unload_start:
        unload_start = _next_day_if_before(arrived, unload_start)

    return Job(
        job_id=job_id,
        source=Source.WORK_LOG,
        ticket_no=row.get("ticket_no") or None,
        load_time=datetime.combine(day, _time(row.get("time")) or time(12, 0)),
        vehicle_reg=row["reg"],
        driver=row.get("driver") or None,
        customer=row["customer"],
        leg=Leg.RETURN if kind == "return" else Leg.DELIVERY,
        carries_load=kind in {"load", "return"},
        from_site=row.get("from_site") or None,
        to_site=row.get("to_site") or None,
        material=row.get("material") or None,
        recycled=_flag(row, "recycled"),
        net_t=_decimal(row.get("net_t")),
        miles=_decimal(row.get("miles")),
        full_return=_flag(row, "full_return"),
        daywork=daywork,
        arrived=arrived,
        unload_start=unload_start,
        internal_account=_flag(row, "internal"),
        source_file=name,
    )


def _daywork(row: dict[str, str], day: date) -> Daywork:
    """Same rule as a ticket: chute pays the written hours, else the whole time on site."""
    chute = _flag(row, "chute")
    written = _decimal(row.get("daywork_hours"))
    on_site = None
    time_in, time_out = _moment(day, row.get("time_in")), _moment(day, row.get("time_out"))
    if time_in and time_out:
        time_out = _next_day_if_before(time_in, time_out)
        on_site = Decimal((time_out - time_in).total_seconds()) / 3600
    hours = written if chute and written is not None else on_site if on_site is not None else written
    if hours is None:
        raise WorkLogError("daywork row needs time_in and time_out, or daywork_hours")
    return Daywork(chute=chute, hours=hours)


def _cell(v) -> str:
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, (date, time)):
        return v.isoformat()
    return str(v)


def _flag(row: dict[str, str], col: str) -> bool:
    return row.get(col, "").lower() in TRUE


def _decimal(v: str | None) -> Decimal | None:
    return Decimal(v.replace(",", "")) if v else None


def _date(v: str) -> date:
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            pass
    raise WorkLogError(f"can't read date {v!r}")


def _time(v: str | None) -> time | None:
    if not v:
        return None
    for fmt in ("%H:%M", "%H:%M:%S", "%H.%M"):
        try:
            return datetime.strptime(v, fmt).time()
        except ValueError:
            pass
    raise WorkLogError(f"can't read time {v!r}")


def _moment(day: date, v: str | None) -> datetime | None:
    if not v:
        return None
    if len(v) > 8:
        return datetime.fromisoformat(v)
    return datetime.combine(day, _time(v))


def _next_day_if_before(start: datetime, end: datetime) -> datetime:
    return end + timedelta(days=1) if end < start else end
